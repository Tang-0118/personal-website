# -*- coding: utf-8 -*-
"""由网页版数据生成 Word 内容母版 resume.docx（第一次或需要重置时跑）。

之后就以 resume.docx 为准：在 Word 里改文字 -> 跑 make_pdf.ps1 ->
cv_from_docx.py 会把 Word 内容读回网页版，再导出网页排版的 PDF。

用法：python build_master.py [--out resume.docx]
"""

import argparse
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

import build_cv_html as cv  # 网页版的数据与渲染器

HERE = os.path.dirname(os.path.abspath(__file__))
FONT = "微软雅黑"
C_TEXT = RGBColor(0x19, 0x23, 0x2D)
C_GRAY = RGBColor(0x5A, 0x64, 0x70)
C_ACCENT = RGBColor(0x5C, 0x4A, 0xDC)
STRONG_RE = re.compile(r"<strong>(.*?)</strong>")


def plain(text: str) -> str:
    """去掉渲染器用的 <strong> 标记，保留文字。"""
    return STRONG_RE.sub(r"\1", text)


def para(doc, text="", size=10, bold=False, color=C_TEXT, space_before=0, space_after=2,
         line=1.15, style=None, keep_next=False):
    p = doc.add_paragraph(style=style)
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    if keep_next:
        pf.keep_with_next = True
    if text:
        run = p.add_run(text)
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run._element.rPr.rFonts.set(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", FONT
        )
    return p


def bullet(doc, head="", text=""):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.15
    pf.left_indent = Pt(16)
    pf.first_line_indent = Pt(-16)
    for chunk, bold in (("•\t", False), (head, True), (text, False)):
        if not chunk:
            continue
        run = p.add_run(chunk)
        run.font.name = FONT
        run.font.size = Pt(10)
        run.font.bold = bold
        run.font.color.rgb = C_TEXT
        run._element.rPr.rFonts.set(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", FONT
        )
    return p


def rich_para(doc, markup: str, size=10, space_after=1):
    """把带 <strong> 的内容写成 Word 段落（加粗部分在 Word 里就是粗体）。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.15
    for chunk in re.split(r"(<strong>.*?</strong>)", markup):
        if not chunk:
            continue
        bold = chunk.startswith("<strong>")
        run = p.add_run(plain(chunk))
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = C_TEXT
        run._element.rPr.rFonts.set(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", FONT
        )
    return p


def section(doc, title):
    p = para(doc, f"[{title}]", size=12, bold=True, space_before=10, space_after=4, keep_next=True)
    return p


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(HERE, "resume.docx"))
    args = ap.parse_args()

    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Pt(36)
        s.left_margin = s.right_margin = Pt(48)

    # ── 抬头 ──
    para(doc, "陈可欣", size=22, bold=True, space_after=2)
    para(doc, "求职意向：计算机视觉 / 图形图像算法（实习）", size=10.5, color=C_GRAY, space_after=2)
    para(doc, plain("专注计算机视觉与深度学习算法的工程落地：目标检测、图像量测与工业质检。"),
         size=10.5, space_after=2)
    para(doc, "联系方式：15319902705（同微信） · 2753538923@qq.com · GitHub: Tang-0118",
         size=10, color=C_GRAY, space_after=2)

    # ── 教育背景 ──
    section(doc, "教育背景")
    para(doc, "湖南大学\t2024.09 — 2028.06", size=11, bold=True, keep_next=True)
    para(doc, "计算机学院 · 智能科学与技术（本科在读）")
    para(doc, "GPA：3.46 / 4.0（86.61 / 100）")
    para(doc, "主修课程：深度学习、数字图像处理、机器学习、模式识别、线性代数、离散数学、程序设计（C++）、"
              "数据结构、操作系统、计算机体系结构、计算机网络、计算机系统等")

    # ── 技术技能 ──
    section(doc, "技术技能")
    for _icon, label, chips in cv.SKILLS:
        para(doc, f"{label}：{'、'.join(chips)}", space_after=1)

    # ── 获奖与认证 ──
    section(doc, "获奖与认证")
    for item in cv.AWARDS:
        rich_para(doc, item)

    # ── 项目经历 ──
    section(doc, "项目经历")
    for proj in cv.PROJECTS:
        head = f"{proj['title']}（{proj['tag']}）" if proj["tag"] else proj["title"]
        para(doc, f"{head}\t{proj['right']}", size=11, bold=True, space_before=6, keep_next=True)
        para(doc, f"一句话：{plain(proj['desc'])}", space_after=1)
        para(doc, f"概述：{plain(proj['summary'])}", space_after=1)
        if proj.get("award"):
            para(doc, f"获奖：{plain(proj['award'])}", space_after=1)
        para(doc, f"技术栈：{'、'.join(proj['chips'])}", size=9.5, color=C_ACCENT, space_after=2)
        for h, t in proj["bullets"]:
            bullet(doc, h, plain(t))

    # ── 实践经历 ──
    section(doc, "实践经历")
    for entry in cv.PRACTICE:
        para(doc, f"{entry['title']}\t{entry['date']}", size=11, bold=True, space_before=6, keep_next=True)
        para(doc, entry["sub"], color=C_GRAY, space_after=1)
        for b in entry["bullets"]:
            bullet(doc, "", plain(b))

    doc.save(args.out)
    print("saved", args.out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
